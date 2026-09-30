# core/excel_exporter.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_assessment(metadata, checks, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    norm_font = Font(name="Calibri", size=10)
    sub_font = Font(name="Calibri", size=12, bold=True, color="333333")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    pass_fill = PatternFill(start_color="D4EDDA", fill_type="solid")
    fail_fill = PatternFill(start_color="F8D7DA", fill_type="solid")
    warn_fill = PatternFill(start_color="FFF3CD", fill_type="solid")
    info_fill = PatternFill(start_color="D1ECF1", fill_type="solid")

    # -------------------------------------------------------------
    # Hoja 1: INFORMACIÓN DEL DISPOSITIVO
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="INFORMACIÓN DEL DISPOSITIVO")
    ws1.views.sheetView[0].showGridLines = True

    total = len(checks)
    pass_cnt = sum(1 for c in checks if c.get('estado') == 'PASS')
    fail_cnt = sum(1 for c in checks if c.get('estado') == 'FAIL')
    warn_cnt = sum(1 for c in checks if c.get('estado') == 'WARNING')
    info_cnt = sum(1 for c in checks if c.get('estado') == 'INFO')

    crit_cnt = sum(1 for c in checks if c.get('sev') == 'Crítica' and c.get('estado') in ['FAIL', 'WARNING'])
    high_cnt = sum(1 for c in checks if c.get('sev') == 'Alta' and c.get('estado') in ['FAIL', 'WARNING'])
    med_cnt = sum(1 for c in checks if c.get('sev') == 'Media' and c.get('estado') in ['FAIL', 'WARNING'])
    low_cnt = sum(1 for c in checks if c.get('sev') == 'Baja' and c.get('estado') in ['FAIL', 'WARNING'])

    # Cálculo ponderado: solo checks con peso aplicable > 0
    app_w = sum(c.get('peso', 0) for c in checks)
    warn_w = sum(c.get('peso', 0) for c in checks if c.get('estado') == 'WARNING')
    earned_w = sum(c.get('peso', 0) for c in checks if c.get('estado') == 'PASS') + (warn_w * 0.5)
    score_global = round((earned_w / app_w) * 100, 1) if app_w > 0 else 0

    ws1['A1'] = "FORTIGATE HARDENING ASSESSMENT"
    ws1['A1'].font = Font(name="Calibri", size=16, bold=True, color="1F497D")

    dashboard_rows = [
        ("INFORMACIÓN DEL DISPOSITIVO", 4, [
            ("Hostname", metadata.get('Hostname', 'N/A')),
            ("Serial", metadata.get('Serial', 'N/A')),
            ("Firmware", metadata.get('Firmware', 'N/A')),
            ("VDOM(s)", metadata.get('VDOM', 'root')),
            ("Fecha de análisis", metadata.get('Fecha', ''))
        ]),
        ("RESULTADO GLOBAL", 11, [
            ("Score Global", f"{score_global}%"),
            ("Madurez", "BÁSICA" if score_global < 70 else "INTERMEDIA" if score_global < 90 else "AVANZADA")
        ]),
        ("ESTADÍSTICAS GLOBALES", 15, [
            ("Total Checks", total),
            ("PASS", pass_cnt),
            ("FAIL", fail_cnt),
            ("WARNING", warn_cnt),
            ("INFO", info_cnt),
            ("CRITICAL", crit_cnt),
            ("HIGH", high_cnt),
            ("MEDIUM", med_cnt),
            ("LOW", low_cnt)
        ])
    ]

    for title, start_row, items in dashboard_rows:
        ws1.cell(row=start_row, column=1, value=title).font = sub_font
        curr = start_row + 1
        for k, v in items:
            c1 = ws1.cell(row=curr, column=1, value=k)
            c2 = ws1.cell(row=curr, column=2, value=v)
            c1.font = bold_font
            c2.font = norm_font
            c1.border = thin_border
            c2.border = thin_border
            curr += 1

    # -------------------------------------------------------------
    # Hoja 2: Resumen por Categoría
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Resumen por Categoría")
    ws2.views.sheetView[0].showGridLines = True
    h2 = ["Categoría", "Score", "PASS", "FAIL", "WARNING", "SKIP", "INFO", "Total", "Peso Aplicable", "Peso Fallido", "Peso WARNING"]
    for idx, h in enumerate(h2, 1):
        c = ws2.cell(row=1, column=idx, value=h)
        c.fill, c.font = header_fill, header_font
        c.alignment = Alignment(horizontal="center")

    cats = {}
    for c in checks:
        cat = c.get('cat', 'General')
        estado = c.get('estado', 'INFO')
        peso = c.get('peso', 0)

        if cat not in cats:
            cats[cat] = {'PASS': 0, 'FAIL': 0, 'WARNING': 0, 'SKIP': 0, 'INFO': 0, 'app': 0, 'fail': 0, 'warn': 0}

        cats[cat][estado] = cats[cat].get(estado, 0) + 1
        cats[cat]['app'] += peso
        if estado == 'FAIL':
            cats[cat]['fail'] += peso
        elif estado == 'WARNING':
            cats[cat]['warn'] += peso

    r = 2
    for cat, st in cats.items():
        earned = st['app'] - st['fail'] - (st['warn'] * 0.5)
        sc = round((earned / st['app']) * 100, 1) if st['app'] > 0 else 100.0
        tot_cat = st['PASS'] + st['FAIL'] + st['WARNING'] + st['SKIP'] + st['INFO']

        ws2.cell(row=r, column=1, value=cat).font = bold_font
        ws2.cell(row=r, column=2, value=f"{max(0, sc)}%").alignment = Alignment(horizontal="right")
        ws2.cell(row=r, column=3, value=st['PASS']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=4, value=st['FAIL']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=5, value=st['WARNING']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=6, value=st['SKIP']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=7, value=st['INFO']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=8, value=tot_cat).alignment = Alignment(horizontal="center")
        ws2.cell(row=r, column=9, value=st['app']).alignment = Alignment(horizontal="right")
        ws2.cell(row=r, column=10, value=st['fail']).alignment = Alignment(horizontal="right")
        ws2.cell(row=r, column=11, value=st['warn']).alignment = Alignment(horizontal="right")
        for col in range(1, 12):
            ws2.cell(row=r, column=col).border = thin_border
        r += 1

    # -------------------------------------------------------------
    # Hoja 3: Detalle de Hallazgos (Con Columna VDOM)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Detalle de Hallazgos")
    ws3.views.sheetView[0].showGridLines = True
    h3 = ["Check ID", "VDOM", "Categoría", "Hallazgo", "Estado", "Severidad", "Peso", "Afectados", "Evidencia", "Recomendación", "Remediación", "Tipo", "Aplicable", "Referencia"]
    for idx, h in enumerate(h3, 1):
        c = ws3.cell(row=1, column=idx, value=h)
        c.fill, c.font = header_fill, header_font
        c.alignment = Alignment(horizontal="center")

    r = 2
    for c in checks:
        ws3.cell(row=r, column=1, value=c.get('id', 'N/A'))
        ws3.cell(row=r, column=2, value=c.get('vdom', 'root')).alignment = Alignment(horizontal="center")
        ws3.cell(row=r, column=3, value=c.get('cat', 'General'))
        ws3.cell(row=r, column=4, value=c.get('hallazgo', ''))
        st_cell = ws3.cell(row=r, column=5, value=c.get('estado', 'INFO'))
        st_cell.alignment = Alignment(horizontal="center")

        if c.get('estado') == 'PASS': st_cell.fill = pass_fill
        elif c.get('estado') == 'FAIL': st_cell.fill = fail_fill
        elif c.get('estado') == 'WARNING': st_cell.fill = warn_fill
        elif c.get('estado') == 'INFO': st_cell.fill = info_fill

        ws3.cell(row=r, column=6, value=c.get('sev', 'Baja')).alignment = Alignment(horizontal="center")
        ws3.cell(row=r, column=7, value=c.get('peso', 0)).alignment = Alignment(horizontal="right")
        ws3.cell(row=r, column=8, value=c.get('afectados', ''))
        ws3.cell(row=r, column=9, value=c.get('evidencia', ''))
        ws3.cell(row=r, column=10, value=c.get('rec', ''))
        ws3.cell(row=r, column=11, value=c.get('rem', ''))
        ws3.cell(row=r, column=12, value=c.get('tipo', 'Técnico'))
        ws3.cell(row=r, column=13, value="Sí" if c.get('peso', 0) > 0 else "No").alignment = Alignment(horizontal="center")
        ws3.cell(row=r, column=14, value=c.get('ref', ''))
        for col in range(1, 15):
            ws3.cell(row=r, column=col).border = thin_border
        r += 1

    # -------------------------------------------------------------
    # Hoja 4: Plan de Remediación Priorizado
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Plan de Remediación Priorizado")
    ws4.views.sheetView[0].showGridLines = True
    h4 = ["Prioridad", "Check ID", "VDOM", "Categoría", "Hallazgo", "Severidad", "Evidencia", "Recomendación", "Remediación", "Referencia"]
    for idx, h in enumerate(h4, 1):
        c = ws4.cell(row=1, column=idx, value=h)
        c.fill, c.font = header_fill, header_font
        c.alignment = Alignment(horizontal="center")

    actionable = [c for c in checks if c.get('estado') in ['FAIL', 'WARNING']]
    sev_map = {"Crítica": 1, "Alta": 2, "Media": 3, "Baja": 4}
    actionable.sort(key=lambda x: sev_map.get(x.get('sev'), 99))

    r = 2
    for c in actionable:
        prio = "Alta" if c.get('sev') in ["Crítica", "Alta"] else "Media" if c.get('sev') == "Media" else "Baja"
        p_cell = ws4.cell(row=r, column=1, value=prio)
        p_cell.alignment = Alignment(horizontal="center")
        p_cell.fill = fail_fill if prio == "Alta" else warn_fill

        ws4.cell(row=r, column=2, value=c.get('id', 'N/A'))
        ws4.cell(row=r, column=3, value=c.get('vdom', 'root')).alignment = Alignment(horizontal="center")
        ws4.cell(row=r, column=4, value=c.get('cat', 'General'))
        ws4.cell(row=r, column=5, value=c.get('hallazgo', ''))
        ws4.cell(row=r, column=6, value=c.get('sev', 'Baja')).alignment = Alignment(horizontal="center")
        ws4.cell(row=r, column=7, value=c.get('evidencia', ''))
        ws4.cell(row=r, column=8, value=c.get('rec', ''))
        ws4.cell(row=r, column=9, value=c.get('rem', ''))
        ws4.cell(row=r, column=10, value=c.get('ref', ''))
        for col in range(1, 11):
            ws4.cell(row=r, column=col).border = thin_border
        r += 1

    # -------------------------------------------------------------
    # Hoja 5: Métricas por Módulo
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Métricas por Módulo")
    ws5.views.sheetView[0].showGridLines = True
    h5 = ["Módulo", "PASS", "FAIL", "WARNING", "SKIP", "INFO", "Total", "Score"]
    for idx, h in enumerate(h5, 1):
        c = ws5.cell(row=1, column=idx, value=h)
        c.fill, c.font = header_fill, header_font
        c.alignment = Alignment(horizontal="center")

    mods = {}
    for c in checks:
        m = c.get('mod', 'General')
        estado = c.get('estado', 'INFO')
        if m not in mods:
            mods[m] = {'PASS': 0, 'FAIL': 0, 'WARNING': 0, 'SKIP': 0, 'INFO': 0}
        mods[m][estado] = mods[m].get(estado, 0) + 1

    r = 2
    for m, st in mods.items():
        tot = st['PASS'] + st['FAIL'] + st['WARNING'] + st['SKIP'] + st['INFO']
        eval_tot = st['PASS'] + st['FAIL'] + st['WARNING']
        sc = round(((st['PASS'] * 10 + st['WARNING'] * 5) / (eval_tot * 10)) * 100, 1) if eval_tot > 0 else 100.0

        ws5.cell(row=r, column=1, value=m).font = bold_font
        ws5.cell(row=r, column=2, value=st['PASS']).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=3, value=st['FAIL']).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=4, value=st['WARNING']).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=5, value=st['SKIP']).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=6, value=st['INFO']).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=7, value=tot).alignment = Alignment(horizontal="center")
        ws5.cell(row=r, column=8, value=f"{sc}%").alignment = Alignment(horizontal="right")
        for col in range(1, 9):
            ws5.cell(row=r, column=col).border = thin_border
        r += 1

    # Autoajuste de ancho de columnas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    wb.save(output_path)
